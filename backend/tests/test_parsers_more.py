"""文本与 XLSX 解析测试：纯代码、粘贴表格、前导零格式、公式缓存与歧义选择。"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from dailyscreen_lite.domain.errors import AmbiguousSheets, ParseFailed
from dailyscreen_lite.domain.models import ParseOptions
from dailyscreen_lite.imports.parsers.text_parser import parse_text
from dailyscreen_lite.imports.parsers.csv_parser import parse_csv
from dailyscreen_lite.imports.parsers.xlsx_parser import parse_xlsx


def test_pure_code_text_is_not_table_without_header():
    result = parse_text("000001\n600519\n300750\n")
    assert result.structure["input_mode"] == "pure_codes"
    assert [c.normalized_code for c in result.candidates] == ["000001", "600519", "300750"]


def test_pure_code_text_rejects_non_code_line_instead_of_scanning_numbers():
    with pytest.raises(ParseFailed):
        parse_text("000001\n股票名称\n600519\n")


def test_pasted_table_with_tab_delimiter_selects_code_column():
    result = parse_text("股票代码\t名称\n000001\t平安银行\n600519\t贵州茅台\n")
    assert [c.normalized_code for c in result.candidates] == ["000001", "600519"]
    assert result.candidates[0].raw_extras["名称"] == "平安银行"


def test_pasted_table_with_whitespace_delimiter():
    result = parse_text("代码 名称\n000001 平安银行\n")
    assert [c.normalized_code for c in result.candidates] == ["000001"]


def test_amount_column_matching_securities_is_not_auto_imported():
    """成交额列内容恰好命中证券库也不能被自动当作代码列，需用户显式选择。"""
    from dailyscreen_lite.domain.errors import AmbiguousColumns

    with pytest.raises(AmbiguousColumns) as excinfo:
        parse_text("名称\t成交额\t序号\nalpha\t600000\t1\nbeta\t600519\t2\n")
    assert excinfo.value.kind == "column"
    # 提示说明列名非代码语义
    assert "列名" in excinfo.value.message


def test_amount_column_can_be_chosen_explicitly_by_user():
    result = parse_text(
        "名称\t成交额\t序号\nalpha\t600000\t1\nbeta\t600519\t2\n",
        options=ParseOptions(code_column=1),
    )
    assert [c.normalized_code for c in result.candidates] == ["600000", "600519"]


def test_text_with_two_code_columns_requires_selection():
    from dailyscreen_lite.domain.errors import AmbiguousColumns

    with pytest.raises(AmbiguousColumns) as excinfo:
        parse_text("a\tb\n000001\t600000\n000002\t600519\n")
    assert excinfo.value.kind == "column"
    assert len(excinfo.value.options) == 2


def _xlsx(rounds: list[tuple[object, str | None]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "候选"
    for value, fmt in rounds:
        sheet.append([value])
        if fmt:
            sheet.cell(row=sheet.max_row, column=1).number_format = fmt
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_xlsx_numeric_code_with_zero_format_preserves_leading_zeros():
    """数值 1 配 000000 显示格式应还原为 000001，而不是当作 1。"""
    data = _xlsx([("代码", None), (1, "000000"), (600519, "000000")])
    result = parse_xlsx(data)
    assert [c.raw_code for c in result.candidates] == ["000001", "600519"]


def test_xlsx_source_credit_footer_is_not_an_unknown_stock():
    data = _xlsx([
        ("股票代码", None),
        ("000001", None),
        ("数据来源于：i问财网站（iwencai.com）", None),
    ])
    result = parse_xlsx(data)
    assert [c.raw_code for c in result.candidates] == ["000001"]
    assert result.declared_total == 1


def test_csv_source_credit_footer_is_ignored_but_invalid_data_row_is_kept():
    result = parse_csv(
        "股票代码,名称\n000001,平安银行\n无效代码,待核查\n数据来源于: i问财网站 (iwencai.com),\n".encode()
    )
    assert [c.raw_code for c in result.candidates] == ["000001", "无效代码"]
    assert result.declared_total == 2


def test_xlsx_multi_sheet_requires_selection_then_resolves():
    workbook = Workbook()
    first = workbook.active
    first.title = "汇总"
    first.append(["说明"])
    first.append(["本次候选在第二个表"])
    second = workbook.create_sheet("候选")
    second.append(["代码"])
    second.append(["000001"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()

    with pytest.raises(AmbiguousSheets) as excinfo:
        parse_xlsx(data)
    assert excinfo.value.kind == "sheet"
    assert {o.value for o in excinfo.value.options} == {"汇总", "候选"}

    result = parse_xlsx(data, options=ParseOptions(sheet="候选"))
    assert [c.normalized_code for c in result.candidates] == ["000001"]


def test_xlsx_formula_without_cache_is_reported_not_silently_dropped():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["代码"])
    sheet.append(["=1+0"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(ParseFailed, match="公式"):
        parse_xlsx(buffer.getvalue())


def test_xlsx_price_column_not_selected_as_code():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["名称", "金额"])
    sheet.append(["alpha", 11.5])
    sheet.append(["beta", 13.2])
    buffer = io.BytesIO()
    workbook.save(buffer)
    # 单列? 不是；两列都不满足代码比例 → 明确失败，不猜
    with pytest.raises(ParseFailed):
        parse_xlsx(buffer.getvalue())
